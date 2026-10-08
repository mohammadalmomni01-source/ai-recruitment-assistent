import os
import re
from pypdf import PdfReader
from docx import Document


def clean_text(text):
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def extract_pdf(path):
    reader = PdfReader(path)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def extract_docx(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def extract_text(path):
    """Returns (text, error). error is None when extraction works."""
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            raw = extract_pdf(path)
        elif ext == ".docx":
            raw = extract_docx(path)
        else:
            return "", f"Unsupported file type: {ext}"
    except Exception as e:
        return "", f"Could not read file: {e}"
    text = clean_text(raw)
    if not text:
        return "", "No text found (file may be empty or a scanned image)"
    return text, None


def extract_cv(path, sender="", subject="", msg_id=""):
    text, error = extract_text(path)
    return {
        "message_id": msg_id,
        "sender": sender,
        "subject": subject,
        "file_path": path,
        "text": text,
        "error": error,
    }


if __name__ == "__main__":
    folder = "temp"
    for name in sorted(os.listdir(folder)):
        result = extract_cv(os.path.join(folder, name))
        print("=" * 40)
        print("File:", name)
        if result["error"]:
            print("ERROR:", result["error"])
        else:
            print("Characters:", len(result["text"]))
            print(result["text"][:300])